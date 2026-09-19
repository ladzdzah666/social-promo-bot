import os
import pandas as pd
from datetime import datetime
from pathlib import Path
import config

class DataManager:
    """Manages reading and updating promo post data from Excel."""

    def __init__(self, file_path: Path = config.DEFAULT_EXCEL_PATH):
        self.file_path = Path(file_path)

    def get_pending_posts(self, check_schedule: bool = True):
        """
        Returns list of posts ready to be published.
        If check_schedule is True, posts whose JADWAL_POSTING is still in the future will be skipped.
        """
        if not self.file_path.exists():
            return []

        # Read first sheet
        df = pd.read_excel(self.file_path, sheet_name=0)
        df.columns = [str(col).strip().upper() for col in df.columns]

        if "STATUS" not in df.columns:
            return []

        df = df.fillna("")
        now = datetime.now()
        ready_posts = []

        for _, row in df.iterrows():
            status = str(row.get("STATUS", "")).strip().lower()
            if status != "pending":
                continue

            post_id = row.get("ID", "")
            caption = str(row.get("CAPTION", "")).strip()
            image_path = str(row.get("IMAGE", row.get("IMAGE_PATH", row.get("GAMBAR", "")))).strip()
            platform = str(row.get("PLATFORM", "both")).strip().lower()
            jadwal_raw = str(row.get("SCHEDULE", row.get("SCHEDULED_TIME", row.get("JADWAL_POSTING", "")))).strip()

            # Schedule checking
            is_ready = True
            scheduled_display = ""
            if check_schedule and jadwal_raw:
                try:
                    # Clean possible timestamp string
                    clean_jadwal = jadwal_raw.split(".")[0]
                    # Try various common formats
                    dt = None
                    for fmt in ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M"]:
                        try:
                            dt = datetime.strptime(clean_jadwal, fmt)
                            break
                        except ValueError:
                            continue

                    if dt:
                        scheduled_display = dt.strftime("%Y-%m-%d %H:%M")
                        if now < dt:
                            # Scheduled in future, skip for now
                            is_ready = False
                except Exception:
                    # If parse fails, assume ready
                    is_ready = True

            if is_ready:
                ready_posts.append({
                    "id": post_id,
                    "caption": caption,
                    "image_path": image_path,
                    "platform": platform,
                    "jadwal": scheduled_display or "Immediate",
                    "status": "pending"
                })

        return ready_posts

    def get_all_pending_summary(self):
        """Returns all posts with status == 'pending' regardless of schedule (for overview)."""
        return self.get_pending_posts(check_schedule=False)

    def mark_post_status(self, post_id, status: str, note: str = ""):
        """Updates the status and timestamp of a post by ID in the Excel file."""
        import openpyxl
        wb = openpyxl.load_workbook(self.file_path)
        ws = wb.active

        # Find header columns
        headers = [str(ws.cell(row=1, column=c).value or "").strip().upper() for c in range(1, ws.max_column + 1)]
        
        id_col = headers.index("ID") + 1 if "ID" in headers else 1
        status_col = headers.index("STATUS") + 1 if "STATUS" in headers else 6
        
        time_candidates = ["POSTED_AT", "PUBLISHED_AT", "WAKTU_TERBIT"]
        time_col = None
        for c_name in time_candidates:
            if c_name in headers:
                time_col = headers.index(c_name) + 1
                break
        if not time_col:
            time_col = 7

        note_candidates = ["NOTES", "NOTE", "KETERANGAN"]
        note_col = None
        for n_name in note_candidates:
            if n_name in headers:
                note_col = headers.index(n_name) + 1
                break
        if not note_col:
            note_col = 8

        found = False
        for r in range(2, ws.max_row + 1):
            cell_val = str(ws.cell(row=r, column=id_col).value or "").strip()
            if cell_val == str(post_id):
                ws.cell(row=r, column=status_col).value = status
                if status.lower() == "posted":
                    ws.cell(row=r, column=time_col).value = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                if note:
                    existing_note = str(ws.cell(row=r, column=note_col).value or "")
                    ws.cell(row=r, column=note_col).value = f"{existing_note} | {note}".strip(" | ")
                found = True
                break

        if found:
            wb.save(self.file_path)
            print(f"[DataManager] Post #{post_id} status updated -> {status}")
        else:
            print(f"[DataManager] Warning: Post #{post_id} not found in spreadsheet.")
